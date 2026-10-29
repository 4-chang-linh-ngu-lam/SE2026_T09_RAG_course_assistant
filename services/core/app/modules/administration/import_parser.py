"""Bounded CSV/XLSX parsing: no raw uploads or validation input in logs/results."""

import base64
import binascii
import csv
import io
import zipfile
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from app.common.errors import AppError
from app.modules.administration.schemas import ImportErrorResponse, ImportPreviewRequest
from app.modules.identity.public import StudentImportRow

MAX_FILE_BYTES = 2_000_000
MAX_XML_BYTES = 10_000_000
MAX_ROWS = 100
MAX_COLUMNS = 64


def parse_students(request: ImportPreviewRequest) -> tuple[list[dict], list[ImportErrorResponse]]:
    try:
        content = base64.b64decode(request.content_base64, validate=True)
        if len(content) > MAX_FILE_BYTES:
            raise ValueError("File too large")
        source_rows = []
        if request.format == "CSV":
            matrix = []
            reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
            previous_line = 0
            for values in reader:
                matrix.append(values)
                source_rows.append(previous_line + 1)
                previous_line = reader.line_num
                if len(matrix) > MAX_ROWS + 1:
                    raise ValueError("Too many rows")
                if len(values) > MAX_COLUMNS:
                    raise ValueError("Too many columns")
        else:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                infos = archive.infolist()
                if len(infos) > 100 or sum(info.file_size for info in infos) > MAX_XML_BYTES:
                    raise ValueError("Archive limits exceeded")
                for info in infos:
                    if info.filename.endswith(".xml"):
                        xml = archive.read(info).upper()
                        if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                            raise ValueError("Unsafe XML")
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(workbook.worksheets) != 1:
                    raise ValueError("Exactly one worksheet required")
                sheet = workbook.worksheets[0]
                if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
                    raise ValueError("Worksheet limits exceeded")
                matrix = []
                for row_number, cells in enumerate(sheet.iter_rows(), start=1):
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValueError("Formula cells are not accepted")
                    matrix.append([cell.value for cell in cells])
                    source_rows.append(row_number)
                    if len(matrix) > MAX_ROWS + 1:
                        raise ValueError("Too many rows")
            finally:
                workbook.close()
        if not matrix or len(matrix) < 2:
            raise ValueError("Missing rows")
        raw_header = matrix[0]
        if not all(isinstance(value, str) for value in raw_header):
            raise ValueError("Invalid columns")
        header = [value.strip() for value in raw_header]
        if len(header) > MAX_COLUMNS or len({value.casefold() for value in header}) != len(header):
            raise ValueError("Invalid columns")
        if not all(header):
            raise ValueError("Invalid columns")
        display_name_header = request.display_name_header
        if (
            "display_name_header" not in request.model_fields_set
            and display_name_header is None
            and "display_name" in header
        ):
            display_name_header = "display_name"
        selected_columns = [(request.username_header, "username")]
        if display_name_header is not None:
            selected_columns.append((display_name_header, "display_name"))
        if request.user_id_header is not None:
            selected_columns.append((request.user_id_header, "user_id"))
        selected_headers = [source.casefold() for source, _target in selected_columns]
        if any(source not in header for source, _target in selected_columns) or len(
            selected_headers
        ) != len(set(selected_headers)):
            raise ValueError("Invalid column mapping")
        column_indexes = [(header.index(source), target) for source, target in selected_columns]
        rows, errors, names, user_ids = [], [], set(), set()
        for index, values in enumerate(matrix[1:], start=1):
            number = source_rows[index]
            if not any(value is not None and value != "" for value in values):
                continue
            try:
                if len(values) != len(header) or any(
                    values[index] is not None
                    and (not isinstance(values[index], str) or len(values[index]) > 500)
                    for index, _target in column_indexes
                ):
                    raise ValueError("Invalid row")
                data = {
                    target: values[index]
                    for index, target in column_indexes
                    if values[index] not in (None, "")
                }
                row = StudentImportRow.model_validate(data)
                if row.username in names or (row.user_id is not None and row.user_id in user_ids):
                    errors.append(ImportErrorResponse(row=number, code="DUPLICATE_ROW"))
                names.add(row.username)
                if row.user_id is not None:
                    user_ids.add(row.user_id)
                rows.append({"number": number, "values": row.model_dump(mode="json")})
            except (ValueError, ValidationError):
                errors.append(ImportErrorResponse(row=number, code="INVALID_ROW"))
        if not rows and not errors:
            raise ValueError("Missing rows")
        return rows, errors
    except (
        ParseError,
        DefusedXmlException,
        InvalidFileException,
        RuntimeError,
        ValueError,
        binascii.Error,
        csv.Error,
        UnicodeError,
        zipfile.BadZipFile,
        KeyError,
        OSError,
    ) as exc:
        raise AppError(422, "Invalid or oversized student import file") from exc
